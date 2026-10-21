"""Reportes compartidos por la pantalla, Excel y PDF."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.logic.caja import caja_flows
from app.logic.catalog import variant_state
from app.logic.profit import profit_report
from app.logic.sales import days_open, layaway_state
from app.logic.common import setting, setting_int
from app.models import (
    Account,
    AccountMove,
    CashSession,
    Customer,
    Expense,
    Layaway,
    Product,
    Purchase,
    Sale,
    User,
    Variant,
)
from app.utils import as_utc, fmt_date, fmt_dt, money, period_bounds, utcnow

FONT = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
FONT_BOLD = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")

TITLES = {
    "ventas": "Ventas",
    "inventario": "Inventario",
    "cartera": "Cartera",
    "separados": "Separados",
    "compras": "Compras",
    "gastos": "Gastos",
    "cuentas": "Cuentas de dinero",
    "cierres": "Cierres de caja",
    "ganancias": "Ganancias",
}


def build_report(
    db: Session,
    kind: str,
    *,
    start: date,
    end: date,
    category_id: int | None = None,
    customer_id: int | None = None,
    user_id: int | None = None,
    product_id: int | None = None,
) -> dict:
    if kind not in TITLES:
        from app.logic.common import BusinessError

        raise BusinessError("Ese reporte no existe.")
    filters = _filter_label(db, start, end, category_id, customer_id, user_id, product_id)
    headers, rows = _rows(db, kind, start, end, category_id, customer_id, user_id, product_id)
    return {
        "kind": kind,
        "title": f"{setting(db, 'business_name')} · {TITLES[kind]}",
        "filters": filters,
        "generated": fmt_dt(utcnow()),
        "headers": headers,
        "rows": rows,
    }


def to_xlsx(report: dict) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = TITLES.get(report["kind"], "Reporte")[:31]
    bold = Font(bold=True, name="Calibri")
    sheet.append([report["title"]])
    sheet.append([f"Generado: {report['generated']}"])
    sheet.append([f"Filtros: {report['filters']}"])
    sheet.append([])
    sheet.append(report["headers"])
    for cell in sheet[5]:
        cell.font = bold
    for row in report["rows"]:
        sheet.append(row)
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        sheet.column_dimensions[letter].width = width
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def to_pdf(report: dict) -> bytes:
    from fpdf import FPDF

    wide = len(report["headers"]) > 6
    pdf = FPDF(orientation="L" if wide else "P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_font("DejaVu", "", str(FONT))
    pdf.add_font("DejaVu", "B", str(FONT_BOLD))
    pdf.add_page()
    pdf.set_font("DejaVu", "B", 14)
    pdf.cell(0, 8, report["title"], new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 8)
    pdf.multi_cell(0, 4, f"Generado: {report['generated']}\nFiltros: {report['filters']}")
    pdf.ln(2)
    usable = pdf.w - 20
    count = max(len(report["headers"]), 1)
    width = usable / count
    pdf.set_font("DejaVu", "B", 7)
    for header in report["headers"]:
        pdf.cell(width, 6, _clip(header, 28), border=1)
    pdf.ln()
    pdf.set_font("DejaVu", "", 7)
    for row in report["rows"]:
        if pdf.get_y() > pdf.h - 16:
            pdf.add_page()
        for value in row:
            pdf.cell(width, 6, _clip(value, 32), border=1)
        pdf.ln()
    if not report["rows"]:
        pdf.cell(0, 8, "Sin movimientos para estos filtros.", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _clip(value, size: int) -> str:
    text = str(value if value is not None else "")
    return text if len(text) <= size else text[: size - 1] + "…"


def _filter_label(db, start, end, category_id, customer_id, user_id, product_id) -> str:
    bits = [f"del {fmt_date(start)} al {fmt_date(end)}"]
    if category_id:
        category = db.get(Product, 0)
        from app.models import Category

        found = db.get(Category, category_id)
        if found:
            bits.append(f"categoría {found.name}")
    if product_id:
        product = db.get(Product, product_id)
        if product:
            bits.append(f"producto {product.name}")
    if customer_id:
        customer = db.get(Customer, customer_id)
        if customer:
            bits.append(f"clienta {customer.name}")
    if user_id:
        person = db.get(User, user_id)
        if person:
            bits.append(f"usuario {person.full_name}")
    return ", ".join(bits)


def _rows(db, kind, start, end, category_id, customer_id, user_id, product_id):
    opening, closing = period_bounds(start, end)
    if kind == "ventas":
        return _ventas(db, opening, closing, category_id, customer_id, user_id, product_id)
    if kind == "inventario":
        return _inventario(db, category_id, product_id)
    if kind == "cartera":
        return _cartera(db, customer_id)
    if kind == "separados":
        return _separados(db, opening, closing, customer_id)
    if kind == "compras":
        return _compras(db, start, end)
    if kind == "gastos":
        return _gastos(db, start, end, category_id)
    if kind == "cuentas":
        return _cuentas(db, opening, closing)
    if kind == "cierres":
        return _cierres(db, opening, closing, user_id)
    return _ganancias(db, start, end)


def _ventas(db, opening, closing, category_id, customer_id, user_id, product_id):
    headers = ["Fecha", "Número", "Clienta", "Cajero", "Origen", "Estado", "Subtotal", "Descuento", "Total", "Cobrado", "Saldo", "Costo", "Ganancia"]
    sales = db.scalars(select(Sale).where(Sale.created_at >= opening, Sale.created_at <= closing).order_by(Sale.created_at)).all()
    rows = []
    for sale in sales:
        if customer_id and sale.customer_id != customer_id:
            continue
        if user_id and sale.user_id != user_id:
            continue
        if product_id or category_id:
            matched = False
            for line in sale.lines:
                if product_id and line.variant.product_id != product_id:
                    continue
                if category_id and line.variant.product.category_id != category_id:
                    continue
                matched = True
            if not matched:
                continue
        cost = sum(line.unit_cost * (line.qty - line.returned_stock) for line in sale.lines) if sale.status != "anulada" else 0
        revenue = sale.total if sale.status != "anulada" else 0
        rows.append([
            fmt_dt(sale.created_at), sale.number, sale.customer.name if sale.customer else "Mostrador",
            sale.user.full_name, "Separado" if sale.source == "separado" else "Mostrador", sale.status,
            sale.subtotal, sale.discount_amount, sale.total, sale.paid, sale.balance, cost, revenue - cost,
        ])
    return headers, rows


def _inventario(db, category_id, product_id):
    headers = ["SKU", "Producto", "Categoría", "Variante", "Unidad", "Paquetes", "Gramos", "Reservado", "No vendible", "Disponible", "Mínimo", "Costo", "Precio"]
    variants = db.scalars(select(Variant).join(Product).order_by(Product.name, Variant.sku)).all()
    rows = []
    for variant in variants:
        if product_id and variant.product_id != product_id:
            continue
        if category_id and variant.product.category_id != category_id:
            continue
        state = variant_state(variant)
        rows.append([
            variant.sku, variant.product.name, variant.product.category.name, variant.short, variant.unit,
            variant.stock_packages, variant.stock_grams, state["reserved"],
            variant.unsellable_units or variant.unsellable_packages or variant.unsellable_grams,
            state["available"] if variant.unit == "unidad" else f"{state['packages']} pq / {state['grams']} g",
            variant.min_stock, variant.cost, variant.price,
        ])
    return headers, rows


def _cartera(db, customer_id):
    headers = ["Número", "Clienta", "Teléfono", "Fecha", "Vence", "Total", "Saldo", "Estado"]
    sales = db.scalars(select(Sale).where(Sale.balance > 0, Sale.status != "anulada").order_by(Sale.due_date)).all()
    rows = []
    for sale in sales:
        if customer_id and sale.customer_id != customer_id:
            continue
        rows.append([
            sale.number, sale.customer.name if sale.customer else "", sale.customer.phone if sale.customer else "",
            fmt_date(sale.created_at), fmt_date(sale.due_date), sale.total, sale.balance, sale.status,
        ])
    return headers, rows


def _separados(db, opening, closing, customer_id):
    headers = ["Número", "Clienta", "Producto", "Acordado", "Abonado", "Saldo", "Estado", "Días", "Vence", "Retención"]
    rows_db = db.scalars(select(Layaway).order_by(Layaway.created_at.desc())).all()
    alert = setting_int(db, "layaway_alert_days")
    rows = []
    for layaway in rows_db:
        created = as_utc(layaway.created_at)
        if created is None or created < opening or created > closing:
            if layaway.status == "activo":
                pass
            else:
                continue
        if customer_id and layaway.customer_id != customer_id:
            continue
        rows.append([
            layaway.number, layaway.customer.name, layaway.variant.label, layaway.agreed_price, layaway.paid,
            layaway.balance, layaway_state(layaway, alert), days_open(layaway), fmt_dt(layaway.due_at),
            f"{layaway.retention_percent}%",
        ])
    return headers, rows


def _compras(db, start, end):
    headers = ["Número", "Proveedor", "Fecha", "Mercancía", "Flete", "Total", "Abonado", "Saldo", "Estado"]
    purchases = db.scalars(select(Purchase).where(Purchase.ordered_at >= start, Purchase.ordered_at <= end).order_by(Purchase.ordered_at)).all()
    rows = [[
        item.number, item.supplier.name, fmt_date(item.ordered_at), item.merchandise, item.shipping,
        item.total, item.paid, item.balance, item.status,
    ] for item in purchases]
    return headers, rows


def _gastos(db, start, end, category_id):
    headers = ["Fecha", "Categoría", "Descripción", "Valor", "Estado", "Cuenta", "Responsable"]
    expenses = db.scalars(select(Expense).where(Expense.spent_on >= start, Expense.spent_on <= end).order_by(Expense.spent_on)).all()
    rows = []
    for item in expenses:
        if category_id and item.category_id != category_id:
            continue
        rows.append([
            fmt_date(item.spent_on), item.category.name, item.description, item.amount, item.status,
            item.account.name if item.account else "", item.user.full_name if item.user else "",
        ])
    return headers, rows


def _cuentas(db, opening, closing):
    headers = ["Fecha", "Cuenta", "Tipo", "Valor", "Detalle", "Usuario"]
    moves = db.scalars(
        select(AccountMove).where(AccountMove.created_at >= opening, AccountMove.created_at <= closing).order_by(AccountMove.id)
    ).all()
    rows = []
    for move in moves:
        rows.append([
            fmt_dt(move.created_at), move.account.name, move.kind, move.amount,
            ("ANULADO " if move.voided else "") + (move.note or ""),
            move.user.full_name if move.user else "",
        ])
    balances = db.scalars(select(Account).order_by(Account.name)).all()
    for account in balances:
        rows.append(["Saldo actual", account.name, account.kind, account.balance, "", ""])
    return headers, rows


def _cierres(db, opening, closing, user_id):
    headers = ["Número", "Abrió", "Base", "Contado al abrir", "Ventas efectivo", "Separados efectivo", "Fiados efectivo", "Gastos efectivo", "Esperado", "Contado", "Diferencia", "Traslado", "Base siguiente", "Estado"]
    sessions = db.scalars(select(CashSession).where(CashSession.opened_at >= opening, CashSession.opened_at <= closing).order_by(CashSession.opened_at)).all()
    rows = []
    for session in sessions:
        if user_id and session.opened_by_id != user_id and session.closed_by_id != user_id:
            continue
        flows = caja_flows(db, session) if session.status == "cerrada" or session.status == "abierta" else {}
        rows.append([
            session.number, session.opened_by.full_name, session.opening_base, session.counted_open,
            flows.get("sales", 0), flows.get("layaways", 0), flows.get("credits", 0),
            flows.get("expenses", 0) + flows.get("payroll", 0),
            session.expected if session.expected is not None else flows.get("expected", ""),
            session.counted_close if session.counted_close is not None else "",
            session.close_diff if session.close_diff is not None else "",
            session.transfer_amount, session.closing_base if session.closing_base is not None else "",
            session.status,
        ])
    return headers, rows


def _ganancias(db, start, end):
    report = profit_report(db, start, end)
    headers = ["Concepto", "Valor"]
    rows = [
        ["Venta neta (solo lo entregado)", report["net_sales"]],
        ["Descuentos concedidos", report["discounts"]],
        ["Devoluciones", report["returned_revenue"]],
        ["Costo de lo vendido", report["cogs"]],
        ["Ganancia bruta", report["gross"]],
        ["Merma no vendible", report["merma"]],
        ["Retenciones de separados (otro ingreso)", report["retentions"]],
        ["Gastos", report["expenses"]],
        ["Nómina", report["payroll"]],
        ["Ganancia neta estimada", report["net"]],
        ["Anticipos de separados excluidos", report["layaway_advances_excluded"]],
        ["Dinero en cuentas hoy", report["cash_on_hand"]],
        ["Anticipos todavía no entregados", report["committed_advances"]],
    ]
    return headers, rows


def money_rows(report: dict) -> dict:
    """Copia con los números ya en pesos, para la pantalla."""
    shown = []
    for row in report["rows"]:
        shown.append([money(value) if isinstance(value, int) else value for value in row])
    return {**report, "rows": shown}
