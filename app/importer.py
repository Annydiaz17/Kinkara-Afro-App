"""Plantilla Excel de inventario inicial, con vista previa antes de confirmar."""

from __future__ import annotations

import hashlib
import json
import re
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import DATA
from app.logic.common import BusinessError
from app.models import ImportBatch, Product, Variant
from app.utils import norm_code, parse_money

COLUMNS = [
    "codigo", "nombre", "categoria", "descripcion", "devolvible", "retencion_pct",
    "venta_rapida", "sku", "largo_pulgadas", "textura", "color", "origen", "tipo_cabello",
    "unidad", "peso_paquete_g", "precio_paquete", "precio_venta", "costo", "cantidad",
    "gramos_sueltos", "stock_minimo", "reservado_cantidad", "reservado_cliente",
    "reservado_telefono", "reservado_documento", "reservado_abonado", "reservado_medio",
    "reservado_autoriza", "reservado_cuenta",
]

YES = {"si", "sí", "s", "1", "true", "x"}


def template_xlsx() -> bytes:
    book = Workbook()
    help_sheet = book.active
    help_sheet.title = "Instrucciones"
    lines = [
        "Carga de inventario inicial — KINKARA AFRO",
        "Llena la hoja Plantilla. No cambies los nombres de las columnas.",
        "unidad: unidad o gramo.",
        "Si la unidad es gramo, cantidad son paquetes sellados, costo es el costo del paquete,",
        "precio_venta es el precio por gramo y precio_paquete es el precio del paquete cerrado.",
        "peso_paquete_g es obligatorio para el cabello por gramos.",
        "devolvible, venta_rapida y reservado_autoriza: si o no.",
        "Los precios van en pesos, sin centavos. Puedes escribir 180000 o 180.000.",
        "Si una peluca ya está separada, llena reservado_cliente, reservado_abonado y reservado_autoriza.",
        "reservado_medio: efectivo, nequi, daviplata, transferencia, breb o tarjeta.",
        "Una plantilla ya confirmada no se vuelve a cargar sola. Para repetirla hay que escribir CARGAR DE NUEVO.",
    ]
    for line in lines:
        help_sheet.append([line])
    help_sheet.column_dimensions["A"].width = 110
    sheet = book.create_sheet("Plantilla")
    sheet.append(COLUMNS)
    sheet.auto_filter.ref = "A1:AC1"
    sheet.freeze_panes = "A2"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def parse_workbook(content: bytes, db: Session) -> dict:
    file_hash = hashlib.sha256(content).hexdigest()
    try:
        book = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise BusinessError("No pude leer ese archivo. Usa la plantilla Excel (.xlsx).") from exc
    sheet = book["Plantilla"] if "Plantilla" in book.sheetnames else book.active
    header = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if header[: len(COLUMNS)] != COLUMNS:
        raise BusinessError("La plantilla no coincide. Descarga la hoja de nuevo y no cambies la fila de títulos.")
    seen_sku = set()
    names_by_code: dict[str, str] = {}
    rows = []
    for index, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        raw = {COLUMNS[i]: cells[i] if i < len(cells) else None for i in range(len(COLUMNS))}
        if all(value is None or str(value).strip() == "" for value in raw.values()):
            continue
        data, errors = _validate_row(db, raw, seen_sku, names_by_code)
        data["fila"] = index
        rows.append({"fila": index, "data": data, "errors": errors})
    duplicate = db.scalar(select(ImportBatch.id).where(ImportBatch.file_hash == file_hash)) is not None
    return {"hash": file_hash, "rows": rows, "duplicate": duplicate}


def save_preview(token: str, payload: dict) -> None:
    folder = DATA / "imports"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / f"{token}.json").write_text(json.dumps(payload), encoding="utf-8")


def load_preview(token: str) -> dict:
    if not re.fullmatch(r"[a-f0-9]{32}", token or ""):
        raise BusinessError("La vista previa venció. Carga el archivo otra vez.")
    path = DATA / "imports" / f"{token}.json"
    if not path.exists():
        raise BusinessError("La vista previa venció. Carga el archivo otra vez.")
    return json.loads(path.read_text(encoding="utf-8"))


def _validate_row(db: Session, raw: dict, seen_sku: set[str], names_by_code: dict[str, str]) -> tuple[dict, list[str]]:
    errors = []
    data: dict = {"fila": 0}

    def text(key: str) -> str:
        value = raw.get(key)
        return "" if value is None else str(value).strip()

    def integer(key: str, label: str, minimum: int = 0) -> int:
        value = text(key)
        if value == "":
            return 0
        try:
            number = parse_money(value, label, allow_zero=True)
        except ValueError as exc:
            errors.append(str(exc))
            return 0
        if number < minimum:
            errors.append(f"{label} no puede ser negativo.")
        return number

    try:
        data["codigo"] = norm_code(text("codigo"), "El código")
    except ValueError as exc:
        errors.append(str(exc))
        data["codigo"] = text("codigo").upper()
    data["nombre"] = text("nombre")
    if not data["nombre"]:
        errors.append("Falta el nombre.")
    data["categoria"] = text("categoria")
    if not data["categoria"]:
        errors.append("Falta la categoría.")
    data["descripcion"] = text("descripcion")
    data["devolvible"] = text("devolvible").lower() in YES
    data["venta_rapida"] = text("venta_rapida").lower() in YES
    data["retencion_pct"] = integer("retencion_pct", "La retención")
    if data["retencion_pct"] > 100:
        errors.append("La retención no puede pasar de 100.")
    try:
        data["sku"] = norm_code(text("sku"), "El SKU")
    except ValueError as exc:
        errors.append(str(exc))
        data["sku"] = text("sku").upper()
    if data["sku"] in seen_sku:
        errors.append(f"El SKU {data['sku']} está repetido en el archivo.")
    seen_sku.add(data["sku"])
    if data["sku"] and db.scalar(select(Variant.id).where(Variant.sku == data["sku"])):
        errors.append(f"El SKU {data['sku']} ya existe en el sistema.")
    existing = db.scalar(select(Product).where(Product.code == data["codigo"])) if data["codigo"] else None
    if existing and existing.name.strip().lower() != data["nombre"].strip().lower():
        errors.append(f"El código {data['codigo']} ya pertenece a «{existing.name}».")
    previous = names_by_code.get(data["codigo"])
    if previous and previous != data["nombre"].strip().lower():
        errors.append(f"El código {data['codigo']} aparece con dos nombres distintos.")
    if data["codigo"]:
        names_by_code[data["codigo"]] = data["nombre"].strip().lower()
    for key in ("largo_pulgadas", "textura", "color", "origen", "tipo_cabello"):
        data[key] = text(key)
    data["unidad"] = text("unidad").lower()
    if data["unidad"] not in {"unidad", "gramo"}:
        errors.append("La unidad debe ser unidad o gramo.")
    data["peso_paquete_g"] = integer("peso_paquete_g", "El peso del paquete")
    data["precio_paquete"] = integer("precio_paquete", "El precio del paquete")
    data["precio_venta"] = integer("precio_venta", "El precio de venta")
    data["costo"] = integer("costo", "El costo")
    data["cantidad"] = integer("cantidad", "La cantidad")
    data["gramos_sueltos"] = integer("gramos_sueltos", "Los gramos sueltos")
    data["stock_minimo"] = integer("stock_minimo", "El stock mínimo")
    if data["unidad"] == "gramo" and data["peso_paquete_g"] <= 0:
        errors.append("El cabello por gramos necesita el peso fijo del paquete.")
    if data["unidad"] == "gramo" and data["precio_paquete"] <= 0 and data["precio_venta"] <= 0:
        errors.append("Define el precio del paquete o el precio por gramo.")
    data["reservado_cantidad"] = integer("reservado_cantidad", "La cantidad reservada")
    data["reservado_cliente"] = text("reservado_cliente")
    data["reservado_telefono"] = text("reservado_telefono")
    data["reservado_documento"] = text("reservado_documento")
    data["reservado_abonado"] = integer("reservado_abonado", "El abono del separado")
    data["reservado_medio"] = text("reservado_medio").lower() or "efectivo"
    data["reservado_autoriza"] = text("reservado_autoriza").lower() in YES
    data["reservado_cuenta"] = text("reservado_cuenta").lower()
    if data["reservado_cliente"] or data["reservado_cantidad"] or data["reservado_abonado"]:
        if not data["reservado_cliente"]:
            errors.append("El separado necesita el nombre de la clienta.")
        if not data["reservado_autoriza"]:
            errors.append("Falta la autorización de datos de la clienta del separado.")
        reserved_qty = data["reservado_cantidad"] or data["cantidad"]
        if reserved_qty <= 0:
            errors.append("Indica cuántas unidades están reservadas.")
        price = data["precio_venta"] if data["unidad"] == "unidad" else data["precio_paquete"]
        if data["reservado_abonado"] > price * max(reserved_qty, 1):
            errors.append("El abono del separado supera el precio.")
        if data["reservado_medio"] not in {"efectivo", "nequi", "daviplata", "transferencia", "breb", "tarjeta"}:
            errors.append("El medio del abono reservado no es válido.")
    return data, errors
